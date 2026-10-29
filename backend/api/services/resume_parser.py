import re
from io import BytesIO

import pymupdf
from docx import Document


SCANNED_RESUME_ERROR = (
    "This looks like a scanned or empty file. Please upload a text-based PDF or DOCX."
)

_EMAIL = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
_URL = re.compile(r"\b(?:https?://|www\.)[^\s<>]+", re.IGNORECASE)
_BARE_URL = re.compile(
    r"\b(?:[A-Z0-9-]+\.)+(?:com|net|org|io|dev|edu|gov|in|co|me|ai)\b(?:/[^\s<>]*)?",
    re.IGNORECASE,
)
_PHONE = re.compile(
    r"(?<!\w)(?:\+?\d{1,3}[\s.()-]?)?(?:\(?\d{3}\)?[\s.-]?)\d{3}[\s.-]?\d{4}(?!\w)"
)
_ADDRESS_LABEL = re.compile(
    r"(?im)^\s*(?:(?:home|mailing)\s+)?(?:address|location)\s*:\s*[^\r\n]+"
)
_STREET_ADDRESS = re.compile(
    r"(?i)\b\d{1,6}\s+(?:[A-Z0-9.'#-]+\s+){1,5}"
    r"(?:street|st\.?|road|rd\.?|avenue|ave\.?|lane|ln\.?|drive|dr\.?|"
    r"boulevard|blvd\.?|court|ct\.?|way|terrace|place|pl\.?)\b[^\r\n]*"
)


class ResumeProcessingError(ValueError):
    pass


def extract_pdf_text(content: bytes) -> str:
    try:
        with pymupdf.open(stream=content, filetype="pdf") as document:
            return "\n".join(page.get_text("text", sort=True) for page in document)
    except Exception as exc:
        raise ResumeProcessingError("The PDF could not be read. Please upload a valid PDF.") from exc


def extract_docx_text(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ResumeProcessingError("The DOCX could not be read. Please upload a valid DOCX.") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = "\n".join(p.text for p in cell.paragraphs if p.text.strip())
                if cell_text:
                    parts.append(cell_text)
    return "\n".join(parts)


def strip_pii(text: str) -> str:
    text = _ADDRESS_LABEL.sub("[REDACTED]", text)
    text = _STREET_ADDRESS.sub("[REDACTED]", text)
    text = _EMAIL.sub("[REDACTED]", text)
    text = _URL.sub("[REDACTED]", text)
    text = _BARE_URL.sub("[REDACTED]", text)
    return _PHONE.sub("[REDACTED]", text)


def parse_resume_text(content: bytes, extension: str) -> str:
    if extension == ".pdf":
        extracted = extract_pdf_text(content)
    elif extension == ".docx":
        extracted = extract_docx_text(content)
    else:
        raise ResumeProcessingError("Please upload a PDF or DOCX file.")

    if len(extracted.strip()) < 200:
        raise ResumeProcessingError(SCANNED_RESUME_ERROR)
    return strip_pii(extracted)
