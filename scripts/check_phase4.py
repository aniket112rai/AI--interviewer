from __future__ import annotations

import asyncio
import json
import re
import sys
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pymupdf as fitz
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from api.main import app
from common.db import SessionLocal, dispose_database_engine
from common.models import DailyUsage, InterviewSession, Resume, User
from common.resume_store import delete_profile, resume_path
from common.schemas.profile import Profile


TEST_EMAIL = f"phase4-{uuid4().hex}@example.invalid"
TEST_PASSWORD = f"Phase4-{uuid4().hex}!"
CONTACT_DETAILS = (
    "contact@example.invalid",
    "+1 (415) 555-0136",
    "https://candidate.example.com/portfolio",
    "123 Redwood Street, San Francisco, CA 94105",
    "Jordan Sample",
)


def make_pdf(text: str, scanned: bool = False) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_textbox(fitz.Rect(45, 45, 560, 780), text, fontsize=9, lineheight=1.2)
    if scanned:
        pixmap = page.get_pixmap(dpi=120)
        image = pixmap.tobytes("png")
        document.close()
        document = fitz.open()
        page = document.new_page()
        page.insert_image(page.rect, stream=image)
    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Software Engineer Resume", level=1)
    document.add_paragraph("Summary: Backend engineer focused on reliable commerce systems.")
    document.add_paragraph("Skills: Django, Python, AWS, PostgreSQL, Docker")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Experience"
    table.cell(0, 1).text = "Backend Engineer"
    table.cell(1, 0).text = "Example Systems"
    table.cell(1, 1).text = (
        "Built Django services on AWS, improving checkout reliability by 35 percent. "
        "Owned API design, PostgreSQL query tuning, monitoring, and incident response. "
        "Delivered systems handling 20,000 orders each day with retry-safe processing."
    )
    document.add_paragraph("Project: Inventory reconciliation service built with Python and Django.")
    document.add_paragraph("Led a migration that reduced processing time from 42 minutes to 11 minutes.")
    document.add_paragraph("Contact: " + " | ".join(CONTACT_DETAILS))
    document.add_paragraph(
        "Additional engineering details: "
        + "Designed maintainable services, reviewed changes, documented deployment workflows, " * 4
    )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def resume_text() -> str:
    return "\n".join(
        [
            "Jordan Sample | contact@example.invalid | +1 (415) 555-0136",
            "Portfolio: https://candidate.example.com/portfolio",
            "Senior software engineer with seven years building high-volume commerce platforms.",
            "Skills: Python, FastAPI, PostgreSQL, React, AWS, Redis.",
            "Experience: At Northstar Commerce, owned a Python and FastAPI payment service "
            "processing 50,000 requests per day. Added idempotency keys and retry-safe workflows, "
            "reducing duplicate charges by 42 percent and improving p95 latency from 800 ms to 220 ms.",
            "Designed PostgreSQL data models, reviewed React dashboards, and led incident response "
            "for production services. Mentored three engineers and coordinated a gradual migration.",
            "Project: Built a transaction reconciliation platform with Python, FastAPI, PostgreSQL, "
            "and AWS. It reconciled daily settlement files and reduced manual investigation by 60 percent.",
            "Project: Created an internal React dashboard for support and operations teams.",
            "Education: Bachelor of Science in Computer Science.",
            "Address: 123 Redwood Street, San Francisco, CA 94105.",
            "Engineering detail: Designed observable services, documented operational procedures, "
            "and partnered with product and support to improve reliability. " * 3,
        ]
    )


def _assert_profile_is_clean(profile: dict, expected_terms: tuple[str, ...]) -> None:
    assert 3 <= len(profile["probe_targets"]) <= 6, profile
    serialized = json.dumps(profile).lower()
    assert all(term.lower() in serialized for term in expected_terms), serialized
    assert "contact@example.invalid" not in serialized
    assert "415" not in serialized
    assert "candidate.example.com" not in serialized
    assert "123 redwood street" not in serialized
    assert "jordan sample" not in serialized
    assert not re.search(r"\bhttps?://", serialized)


def run_api_check() -> dict:
    pdf_id = ""
    docx_id = ""
    with TestClient(app) as client:
        signup = client.post(
            "/api/auth/signup",
            json={"email": TEST_EMAIL, "password": TEST_PASSWORD, "name": "Phase 4 Check"},
        )
        assert signup.status_code == 201, signup.text
        headers = {"Authorization": f"Bearer {signup.cookies['access_token']}"}

        pdf_upload = client.post(
            "/api/resumes",
            headers=headers,
            files={"file": ("candidate.pdf", make_pdf(resume_text()), "application/pdf")},
        )
        assert pdf_upload.status_code == 202, pdf_upload.text
        pdf_id = pdf_upload.json()["id"]
        pdf_result = client.get(f"/api/resumes/{pdf_id}", headers=headers)
        assert pdf_result.status_code == 200, pdf_result.text
        assert pdf_result.json()["status"] == "ready", pdf_result.text
        _assert_profile_is_clean(pdf_result.json()["profile"], ("python", "fastapi", "50,000"))

        docx_upload = client.post(
            "/api/resumes",
            headers=headers,
            files={"file": ("candidate.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert docx_upload.status_code == 202, docx_upload.text
        docx_id = docx_upload.json()["id"]
        docx_result = client.get(f"/api/resumes/{docx_id}", headers=headers)
        assert docx_result.status_code == 200, docx_result.text
        assert docx_result.json()["status"] == "ready", docx_result.text
        _assert_profile_is_clean(docx_result.json()["profile"], ("django", "aws", "postgresql"))

        scanned_upload = client.post(
            "/api/resumes",
            headers=headers,
            files={"file": ("scanned.pdf", make_pdf(resume_text(), scanned=True), "application/pdf")},
        )
        assert scanned_upload.status_code == 202, scanned_upload.text
        scanned_result = client.get(
            f"/api/resumes/{scanned_upload.json()['id']}", headers=headers
        )
        assert scanned_result.status_code == 200, scanned_result.text
        assert scanned_result.json()["status"] == "failed", scanned_result.text
        assert scanned_result.json()["error"] == (
            "This looks like a scanned or empty file. Please upload a text-based PDF or DOCX."
        ), scanned_result.text

        stranger_signup = client.post(
            "/api/auth/signup",
            json={"email": f"phase4-other-{uuid4().hex}@example.invalid", "password": TEST_PASSWORD},
        )
        assert stranger_signup.status_code == 201
        stranger_headers = {"Authorization": f"Bearer {stranger_signup.cookies['access_token']}"}
        assert client.get(f"/api/resumes/{pdf_id}", headers=stranger_headers).status_code == 404

        snapshot = pdf_result.json()["profile"]
        removed = client.delete(
            f"/api/resumes/{pdf_id}",
            headers=headers,
        )
        assert removed.status_code == 204, removed.text
        assert not resume_path(pdf_id).exists()
        listed = client.get("/api/resumes", headers=headers)
        assert listed.status_code == 200
        assert all(item["id"] != pdf_id for item in listed.json())

    return snapshot


async def cleanup(profile_snapshot: dict | None) -> None:
    async with SessionLocal() as db:
        async with db.begin():
            owner_id = await db.scalar(select(User.id).where(User.email == TEST_EMAIL))
            if owner_id and profile_snapshot:
                resume = Resume(user_id=owner_id, filename="snapshot-check.pdf", status="ready")
                db.add(resume)
                await db.flush()
                interview = InterviewSession(
                    user_id=owner_id,
                    resume_id=resume.id,
                    level="senior",
                    duration_min=15,
                    status="ready",
                    profile_snapshot=profile_snapshot,
                )
                db.add(interview)
                await db.flush()
                await db.delete(resume)
                await db.flush()
                await db.refresh(interview)
                assert interview.resume_id is None
                assert interview.profile_snapshot == profile_snapshot

            users = list(
                (await db.scalars(select(User).where(User.email.like("phase4-%@example.invalid")))).all()
            )
            user_ids = [user.id for user in users]
            resumes = list(
                (await db.scalars(select(Resume).where(Resume.user_id.in_(user_ids)))).all()
            ) if user_ids else []
            resume_ids = [resume.id for resume in resumes]
            for resume_id in resume_ids:
                await delete_profile(resume_id)
            if user_ids:
                await db.execute(delete(DailyUsage).where(DailyUsage.user_id.in_(user_ids)))
                await db.execute(delete(User).where(User.id.in_(user_ids)))
    await dispose_database_engine()


if __name__ == "__main__":
    profile_snapshot = None
    try:
        profile_snapshot = run_api_check()
    finally:
        asyncio.run(cleanup(profile_snapshot))
    print("Phase 4 PDF, DOCX, PII, scanned-file, ownership, delete, and snapshot checks passed.")
